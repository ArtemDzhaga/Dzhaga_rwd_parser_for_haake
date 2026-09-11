"""Temperature/concentration reports, using factual RWD values and nominal T."""
from __future__ import annotations

import math
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import excel_import as source

SHEETS = {
    "ru": ("Сырые данные", "Расчеты", "Итоги", "Легенда"),
    "en": ("Raw data", "Calculations", "Results", "Legend"),
}
HEADERS = {
    "ru": ["Вещество", "Объемная доля, %", "Вал", "Поле", "Номинальная T, °C",
           "Eta, Pa·s", "Изменение с магнитом, %", "Eta(Alum)/Eta(Standart)",
           "Исходный файл", "Строка первой точки", "Строка второй точки"],
    "en": ["Material", "Volume fraction, %", "Shaft", "Field", "Nominal T, °C",
           "Eta, Pa·s", "Magnetic change, %", "Eta(Alum)/Eta(Standart)",
           "Source file", "First point row", "Second point row"],
}


def identity(path):
    name = path.stem
    match = re.search(r"(?:^|_)(Fe3O4|TiO2)_(\d+(?:[.,]\d+)?)", name, re.I)
    if not match:
        raise RuntimeError(f"Material/concentration missing in filename: {path.name}")
    tokens = set(re.split(r"[_\s-]+", name.casefold()))
    shafts = {"Alum" for t in tokens if t in {"custom", "alum"}} | {"Standart" for t in tokens if t in {"standart", "standard"}}
    fields = {"magnet" for t in tokens if t in {"magnit", "magnet"}} | {"control" for t in tokens if t == "control"}
    if len(shafts) != 1 or len(fields) != 1:
        raise RuntimeError(f"Ambiguous or missing shaft/field: {path.name}")
    material = {"fe3o4": "Fe3O4", "tio2": "TiO2"}[match[1].lower()]
    return material, float(match[2].replace(",", ".")), shafts.pop(), fields.pop()


def numeric(value):
    try:
        number = float(value)
        return number if math.isfinite(number) else None
    except (TypeError, ValueError):
        return None


def style_sheet(sheet, widths):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="203650")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 38
    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width


def import_report(inputs, ascii_paths, recursive, temperatures, points, output, language, identities=None):
    if output.exists():
        raise RuntimeError(f"Output already exists: {output}")
    if len(set(temperatures)) != len(temperatures) or not all(math.isfinite(t) for t in temperatures):
        raise RuntimeError("Nominal temperatures must be finite and distinct")
    if len(points) != 2 or points[0] == points[1] or min(points) < 1:
        raise RuntimeError("Select two distinct positive point numbers")
    files, ascii_files = source.collect_csv_source_files(inputs, recursive)
    ascii_files = sorted(set(ascii_files + list(ascii_paths)))
    if not files or not ascii_files:
        raise RuntimeError("RWD files and matching ASCII structures are required")
    structures = [source.extract_ascii_record(p) for p in ascii_files]
    workbook = Workbook()
    workbook.remove(workbook.active)
    raw_name, calc_name, results_name, legend_name = SHEETS[language]
    raw = workbook.create_sheet(raw_name)
    calc = workbook.create_sheet(calc_name)
    workbook.create_sheet(results_name)
    workbook.create_sheet(legend_name)
    raw.append(
        ["Исходный файл", "Вещество", "Объемная доля, %", "Вал", "Поле", "Сегмент",
         "Точка", "Номинальная T, °C", "Eta, Pa·s", "Скорость сдвига, 1/s", "Tau, Pa", "t, s", "t_seg, s"]
        if language == "ru" else
        ["Source file", "Material", "Volume fraction, %", "Shaft", "Field", "Segment",
         "Point", "Nominal T, °C", "Eta, Pa·s", "Shear rate, 1/s", "Tau, Pa", "t, s", "t_seg, s"]
    )
    calc.append(HEADERS[language])
    groups = []
    seen = set()
    for file in files:
        material, concentration, shaft, field = (identities or {}).get(file) or identity(file)
        record = source.extract_rwd_record(file)
        structure, bindings, _ = source.select_ascii_record_for_csv(record, structures)
        if "Eta in Pas" not in structure.headers:
            raise RuntimeError(f"Steady viscosity channel Eta in Pas missing: {file.name}")
        segments = defaultdict(dict)
        for values in source.measurement_rows(record, structure, bindings):
            row = dict(zip(structure.headers, values))
            label = str(values[0])
            match = re.fullmatch(r"(\d+)\|(\d+)", label)
            if not match:
                raise RuntimeError(f"Invalid segment/point label: {label}")
            segment, point = map(int, match.groups())
            if segment < 1 or segment > len(temperatures) or point in segments[segment]:
                raise RuntimeError(f"Unexpected/duplicate segment or point: {file.name}, {label}")
            eta = numeric(row.get("Eta in Pas"))
            raw.append([str(file), material, concentration, shaft, field, segment, point,
                        temperatures[segment - 1], eta, numeric(row.get("GP in 1/s")),
                        numeric(row.get("Tau in Pa")), numeric(row.get("t in s")),
                        numeric(row.get("t_seg in s"))])
            segments[segment][point] = (raw.max_row, eta)
        if set(segments) != set(range(1, len(temperatures) + 1)):
            raise RuntimeError(f"Segment count does not match nominal temperatures: {file.name}")
        for segment, rows in sorted(segments.items()):
            key = (material, concentration, shaft, field, temperatures[segment - 1])
            if key in seen:
                raise RuntimeError(f"Duplicate measurement group: {key}")
            seen.add(key)
            if any(p not in rows or rows[p][1] is None or rows[p][1] <= 0 for p in points):
                raise RuntimeError(f"Missing/invalid Eta at points {points}: {file.name}, segment {segment}")
            groups.append((key, file, rows[points[0]][0], rows[points[1]][0]))
    lookup = {}
    for key, file, first, second in sorted(groups, key=lambda item: item[0]):
        material, concentration, shaft, field, temperature = key
        field_label = {"control": "Контроль", "magnet": "Магнит"}[field] if language == "ru" else field
        calc.append([material, concentration, shaft, field_label, temperature,
                     f"=('{raw_name}'!I{first}+'{raw_name}'!I{second})/2", None, None, str(file), first, second])
        lookup[key] = calc.max_row
    for key, row in lookup.items():
        material, concentration, shaft, field, temperature = key
        control = lookup.get((material, concentration, shaft, "control", temperature))
        standard = lookup.get((material, concentration, "Standart", field, temperature))
        if field == "magnet" and control:
            calc.cell(row, 7, f"=F{row}/F{control}-1")
        if shaft == "Alum" and standard:
            calc.cell(row, 8, f"=F{row}/F{standard}")
        for col in (6, 8):
            calc.cell(row, col).number_format = "0.00000"
        calc.cell(row, 7).number_format = "+0.0%;-0.0%;0.0%"
    style_sheet(raw, [55, 14, 18, 14, 14, 12, 12, 20, 18, 20, 18, 18, 18])
    style_sheet(calc, [15, 20, 14, 14, 22, 18, 24, 27, 55, 18, 18])
    # Row references retain traceability without cluttering the working table.
    calc.column_dimensions["J"].hidden = True
    calc.column_dimensions["K"].hidden = True
    workbook.calculation.fullCalcOnLoad = True
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()


def read_results(path):
    workbook = load_workbook(path, data_only=False)
    for names in SHEETS.values():
        if names[0] in workbook.sheetnames and names[1] in workbook.sheetnames:
            raw, calc = workbook[names[0]], workbook[names[1]]
            break
    else:
        workbook.close()
        raise RuntimeError("Use a B_draft report generated by custom mode")
    results = []
    seen = set()
    for row in calc.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        material, concentration, shaft, field, temperature, eta = row[:6]
        field = {"Контроль": "control", "Магнит": "magnet"}.get(field, field)
        # Evaluate the specific generated formula from its source cells; Excel need not be opened.
        if isinstance(eta, str) and eta.startswith("="):
            match = re.fullmatch(r"=\('([^']+)'!I(\d+)\+'\1'!I(\d+)\)/2", eta)
            if not match or match[1] != raw.title:
                raise RuntimeError("Unrecognized Eta formula; use numeric Eta or the generated two-point mean")
            values = [numeric(raw.cell(int(match[i]), 9).value) for i in (2, 3)]
            eta = sum(values) / 2 if all(v is not None and v > 0 for v in values) else None
        eta = numeric(eta)
        if eta is None or eta <= 0:
            raise RuntimeError(f"Invalid Eta: {material}, {temperature}")
        key = (material, concentration, shaft, field, temperature)
        if key in seen:
            raise RuntimeError(f"Duplicate result: {key}")
        seen.add(key)
        results.append((*key, eta))
    if not results:
        raise RuntimeError("No B_draft measurements found")
    return workbook, names, results


def chart_specs(results, language):
    ru = language == "ru"
    concentration_label = "Объемная доля, %" if ru else "Volume fraction, %"
    temp_label = "Номинальная T, °C" if ru else "Nominal T, °C"
    field_label = {"control": "Контроль", "magnet": "Магнит"} if ru else {"control": "Control", "magnet": "Magnet"}
    specs = []
    for material in sorted({r[0] for r in results}):
        subset = [r for r in results if r[0] == material]
        for shaft, field in sorted({(r[2], r[3]) for r in subset}):
            series = defaultdict(list)
            for _, c, s, f, t, eta in subset:
                if (s, f) == (shaft, field):
                    series[f"{t:g} °C"].append((c, eta))
            specs.append((f"{material}_{shaft}_{field}_concentration", f"{material}: {shaft}, {field_label[field]}", concentration_label, "Eta, Pa·s", series))
        for concentration in sorted({r[1] for r in subset}):
            series = defaultdict(list)
            for _, c, shaft, field, t, eta in subset:
                if c == concentration:
                    series[f"{shaft}, {field_label[field]}"].append((t, eta))
            specs.append((f"{material}_{concentration:g}_temperature", f"{material} {concentration:g}%", temp_label, "Eta, Pa·s", series))
        lookup = {r[:5]: r[5] for r in subset}
        for shaft in sorted({r[2] for r in subset}):
            series = defaultdict(list)
            for _, c, s, field, t, eta in subset:
                control = lookup.get((material, c, s, "control", t))
                if s == shaft and field == "magnet" and control:
                    series[f"{t:g} °C"].append((c, (eta / control - 1) * 100))
            if series:
                ylabel = "Изменение Eta, %" if ru else "Eta change, %"
                specs.append((f"{material}_{shaft}_magnetic_change", f"{material}: {shaft}", concentration_label, ylabel, series))
    return specs


def plot_report(path, output_dir, formats, language, dpi=300):
    workbook, names, results = read_results(path)
    specs = chart_specs(results, language)
    output_dir.mkdir(parents=True, exist_ok=True)
    colors = ["4472C4", "ED7D31", "70AD47", "A64D79", "8064A2"]
    if "xlsx" in formats:
        output = output_dir / f"{path.stem}_charts_{language}.xlsx"
        if output.exists():
            raise RuntimeError(f"Output already exists: {output}")
        sheet = workbook[names[2]]
        if sheet.max_row > 1 or sheet["A1"].value is not None or sheet._charts:
            raise RuntimeError("Results sheet is not empty; select the report produced by the parser")
        data_row = 1
        for index, (_, title, xlabel, ylabel, series) in enumerate(specs):
            chart = ScatterChart()
            chart.scatterStyle = "lineMarker"
            chart.title, chart.x_axis.title, chart.y_axis.title = title, xlabel, ylabel
            chart.width, chart.height = 20, 12
            chart.legend.position = "b"
            for color, (label, points) in enumerate(sorted(series.items())):
                points = sorted(points)
                sheet.cell(data_row, 30, label)
                for offset, (x, y) in enumerate(points, 1):
                    sheet.cell(data_row + offset, 30, x)
                    sheet.cell(data_row + offset, 31, y)
                xref = Reference(sheet, min_col=30, min_row=data_row + 1, max_row=data_row + len(points))
                yref = Reference(sheet, min_col=31, min_row=data_row + 1, max_row=data_row + len(points))
                curve = Series(yref, xref, title=label)
                curve.marker.symbol, curve.marker.size = "circle", 5
                curve.graphicalProperties.line.solidFill = colors[color % len(colors)]
                curve.marker.graphicalProperties.solidFill = colors[color % len(colors)]
                curve.marker.graphicalProperties.line.solidFill = colors[color % len(colors)]
                chart.series.append(curve)
                data_row += len(points) + 2
            sheet.add_chart(chart, f"{'A' if index % 2 == 0 else 'N'}{1 + (index // 2) * 25}")
        sheet.column_dimensions["AD"].width = 30
        sheet.column_dimensions["AE"].width = 18
        workbook.save(output)
    image_formats = [f for f in formats if f != "xlsx"]
    if image_formats:
        from . import plot
        plot.load_plot_dependencies()
        for stem, title, xlabel, ylabel, series in specs:
            fig, ax = plot.plt.subplots(figsize=(8, 5.5), layout="constrained")
            for index, (label, points) in enumerate(sorted(series.items())):
                xs, ys = zip(*sorted(points))
                ax.plot(xs, ys, "-o", label=label, color="#" + colors[index % len(colors)])
            ax.set(title=title, xlabel=xlabel, ylabel=ylabel)
            ax.grid(True, linestyle=":", alpha=0.5)
            ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.14), ncol=2)
            for fmt in image_formats:
                fig.savefig(output_dir / f"{stem}_{language}.{fmt}", dpi=dpi, bbox_inches="tight")
            plot.plt.close(fig)
    workbook.close()
    print(("Графики сохранены: " if language == "ru" else "Plots saved: ") + str(output_dir))
    return 0
