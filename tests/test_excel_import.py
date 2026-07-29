from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

from openpyxl import Workbook

from haake_rheo.excel_import import (
    RwdRecord,
    clear_measurement_region,
    embedded_record_voltage,
    format_frequency_heading,
    normalize_collected_rwd_files,
    record_identity_notes,
    record_experiment_identity,
    rebuild_output_data,
    standardize_frequency,
)


class BatchFileSelectionTests(TestCase):
    def test_keeps_complementary_batch_ranges(self) -> None:
        with TemporaryDirectory() as directory:
            voltage_dir = Path(directory) / "sample" / "500v"
            voltage_dir.mkdir(parents=True)
            low = voltage_dir / "sample_freq=from_01_to_20hz_U=500v_30_points_good.rwd"
            high = voltage_dir / "sample_freq=from_30_to_100hz_U=500v_30_points_good.rwd"
            low.touch()
            high.touch()

            selected = normalize_collected_rwd_files([low, high])

            self.assertEqual({path.name for path in selected}, {low.name, high.name})

    def test_keeps_voltage_from_rwd_filename(self) -> None:
        record = RwdRecord(
            path=Path("/tmp/sample/2000v/sample_freq=from_01_to_100hz_U=2000v_good.rwd"),
            blob=b"",
            size_bytes=0,
            format_magic="",
            rheowin_version="",
            date="",
            time="",
            device="",
            instrument_model="",
            geometry="",
            temperature_controller="",
            gap="",
            job_file="",
            mode="",
            serial_numbers="",
            driver_versions="",
            firmware_versions="",
            strings=[
                (
                    0,
                    r"D:\data\sample\1000v\sample_freq=100hz_U=1000v_30_points.rwd",
                )
            ],
            numeric_blocks=[],
        )

        self.assertEqual(embedded_record_voltage(record), "U=1000v")
        self.assertEqual(record_experiment_identity(record).voltage, "U=2000v")
        self.assertIn("использовано значение из имени файла", record_identity_notes([record])[0])


class WorkbookPreparationTests(TestCase):
    def setUp(self) -> None:
        self.workbook = Workbook()
        self.input_ws = self.workbook.active
        self.input_ws.title = "input_data"
        self.input_ws.cell(1, 2, "sample_U=0v_30_points")
        self.input_ws.cell(2, 2, "Частота 0,1 Гц")
        headers = ["t_seg in s", "Tau in Pa", "G' in Pa", 'G" in Pa', "|eta*| in Pas", "gamma"]
        for offset, header in enumerate(headers):
            self.input_ws.cell(3, 2 + offset, header)
        for point in range(20):
            row = 4 + point
            values = [point + 1, 100 + point, 1000 - point, 100 + point, 2000 - point, 0.001 + point / 1000]
            self.input_ws.cell(row, 1, f"1|{point + 1}")
            for offset, value in enumerate(values):
                self.input_ws.cell(row, 2 + offset, value)

    def test_clears_values_without_removing_layout(self) -> None:
        clear_measurement_region(self.input_ws, 1, self.input_ws.max_row)

        self.assertEqual(self.input_ws.cell(1, 2).value, "sample_U=0v_30_points")
        self.assertEqual(self.input_ws.cell(2, 2).value, "Частота 0,1 Гц")
        self.assertEqual(self.input_ws.cell(3, 4).value, "G' in Pa")
        self.assertIsNone(self.input_ws.cell(4, 4).value)

    def test_rebuilds_output_data_from_actual_series(self) -> None:
        source_count = rebuild_output_data(self.workbook, "input_data", points=20)
        output_ws = self.workbook["output_data"]

        self.assertEqual(source_count, 1)
        self.assertEqual(output_ws.cell(3, 3).value, "sample_0v")
        self.assertEqual(output_ws.cell(5, 3).value, '=IF(\'input_data\'!G4="","",\'input_data\'!G4)')
        self.assertEqual(output_ws.cell(5, 4).value, '=IF(\'input_data\'!D4="","",\'input_data\'!D4)')

    def test_standardizes_nearby_frequency_values(self) -> None:
        self.assertEqual(standardize_frequency(30.66), 30.0)
        self.assertEqual(standardize_frequency(51.1), 50.0)
        self.assertEqual(format_frequency_heading(30.66), "Частота 30 Гц")
        self.assertEqual(format_frequency_heading(51.1), "Частота 50 Гц")
